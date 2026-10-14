"""
Builds sample.pptx, the PowerPoint fixture for the importer tests, with python-pptx's default
template (real layouts, masters and theme). Run from this folder: python3 make-pptx.py
Transitions and animations are injected as raw XML because python-pptx has no API for them.
"""
import io
import struct
import zlib

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches, Pt


def png(w, h, rgb):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 1. Title slide
s1 = prs.slides.add_slide(prs.slide_layouts[0])
s1.shapes.title.text = "Quarterly Review"
s1.placeholders[1].text = "Prepared for Seb"
s1.notes_slide.notes_text_frame.text = "Welcome everyone.\nKeep this short."

# 2. Title and content with bullets and formatting
s2 = prs.slides.add_slide(prs.slide_layouts[1])
s2.shapes.title.text = "What changed"
tf = s2.placeholders[1].text_frame
tf.text = "Revenue grew 12%"
p = tf.add_paragraph()
r = p.add_run()
r.text = "Costs "
r2 = p.add_run()
r2.text = "fell"
r2.font.bold = True
r3 = p.add_run()
r3.text = " slightly"
r3.font.italic = True
p2 = tf.add_paragraph()
p2.level = 1
r4 = p2.add_run()
r4.text = "Details on the website"
r4.hyperlink.address = "https://example.com/report"

# 3. Shapes, picture and a group
s3 = prs.slides.add_slide(prs.slide_layouts[6])
box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1.5))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0x2A, 0x78, 0xD6)
box.text_frame.text = "Inside a shape"
oval = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(1), Inches(2), Inches(2))
oval.fill.solid()
oval.fill.fore_color.rgb = RGBColor(0xEB, 0x68, 0x34)
oval.line.fill.background()
star = s3.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(8), Inches(1), Inches(2), Inches(2))
star.rotation = 15
arrow = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(10.5), Inches(1.5), Inches(2), Inches(1))
line = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(4), Inches(6), Inches(5))
line.line.color.rgb = RGBColor(0xE3, 0x49, 0x48)
line.line.width = Pt(4)
pic = s3.shapes.add_picture(io.BytesIO(png(64, 48, (27, 175, 122))), Inches(7), Inches(4), Inches(2), Inches(1.5))
pic.crop_left = 0.25
grp = s3.shapes.add_group_shape()
g1 = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10), Inches(4), Inches(1), Inches(1))
g2 = grp.shapes.add_textbox(Inches(10), Inches(5.2), Inches(2), Inches(0.6))
g2.text_frame.text = "Grouped"

# 4. Charts and a table
s4 = prs.slides.add_slide(prs.slide_layouts[5])
s4.shapes.title.text = "Numbers"
cd = CategoryChartData()
cd.categories = ["Q1", "Q2", "Q3"]
cd.add_series("North", (10, 14, 18))
cd.add_series("South", (8, 9, 12))
ch = s4.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6), Inches(4), cd).chart
ch.has_title = True
ch.chart_title.text_frame.text = "Sales by region"
cd2 = CategoryChartData()
cd2.categories = ["A", "B", "C", "D"]
cd2.add_series("Share", (40, 30, 20, 10))
s4.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(7), Inches(1.5), Inches(3), Inches(3), cd2)
tbl = s4.shapes.add_table(2, 2, Inches(7), Inches(5), Inches(4), Inches(1)).table
tbl.cell(0, 0).text = "Region"
tbl.cell(0, 1).text = "Total"
tbl.cell(1, 0).text = "North"
tbl.cell(1, 1).text = "42"

# 5. Hidden slide with a dark background and a video
s5 = prs.slides.add_slide(prs.slide_layouts[6])
s5.background.fill.solid()
s5.background.fill.fore_color.rgb = RGBColor(0x14, 0x16, 0x1C)
s5.shapes.add_movie("tiny.mp4", Inches(2), Inches(1), Inches(6), Inches(3.375), poster_frame_image=io.BytesIO(png(32, 18, (40, 40, 40))), mime_type="video/mp4")
s5._element.set("show", "0")

# Transition on slide 2 and a click-to-build animation on its body placeholder.
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
tr = etree.SubElement(s2._element, f"{{{P}}}transition", spd="slow")
etree.SubElement(tr, f"{{{P}}}push", dir="u")
body_id = s2.placeholders[1].shape_id
timing = etree.fromstring(f"""
<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{body_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="7" dur="750"/><p:tgtEl><p:spTgt spid="{body_id}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>""")
s2._element.append(timing)

prs.save("sample.pptx")
print("wrote sample.pptx")
